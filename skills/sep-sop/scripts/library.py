#!/usr/bin/env python3
"""持久词库（哥飞「关键词库」标准的本地版）。

与每日台账的区别：台账是单轮挖掘的流水，词库是跨轮累积的全量库——
每个词只留一行，记录首次/最近出现、生命周期状态、GPTs 倍数、AI 释义、
内外页计数和月趋势（带数据条，替代 sparkline）。扫描次数攒出来的存量
就是这个文件，跑一天攒一天。

状态流转（自动）：
  新词      首次出现 ≤90 天
  二次爆火  峰值在 ≥4 个月前，近期（近3月均值）从峰后谷底回升 ≥2 倍且 ≥ 峰值 25%
  长期大词  出现 ≥120 天且基线（前12月）均值 ≥1000
  观察中    以上都不满足
  已推荐    人工/业务拍板后由 mark 设置（粘性状态，不再降级）
  不值得做  阶段 0.5 被 LLM 过滤的词（新词入库时标注；库内已有词不被降级）

用法:
  python library.py merge --rising /tmp/rising_clean.json [--dropped /tmp/rising_dropped.json] [--monthly /tmp/monthly.json]
  python library.py mark --words "clover fall codes,devara song"
  python library.py stats
  python library.py export --out /tmp/library_new.json      # 导出指定状态的词（默认新词+二次爆火，给挖掘流水线吃）

--monthly 文件格式: {"<词>": {"2025-08": 1234, ...}}（mine_round.py --export-monthly 产出）
词库路径可用 KEYWORD_LIBRARY 环境变量覆盖，默认与台账同目录。
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

LIBRARY = Path(os.environ.get(
    "KEYWORD_LIBRARY",
    "/Users/xg/Projects/SEO-Projects/关键词库.xlsx"))

FIXED_HEADERS = ["词", "词根", "首次出现", "最近更新", "状态", "GPTs倍数", "峰值月",
                 "释义", "内页候选", "专门页面", "推荐日期", "备注"]
STATE_ORDER = {"新词": 0, "二次爆火": 1, "长期大词": 2, "观察中": 3,
               "已推荐": 4, "不值得做": 5}
MONTH_RE = re.compile(r"^\d{4}-\d{2}$")


def _today() -> dt.date:
    return dt.date.today()


def _ym(d: dt.date) -> str:
    return d.strftime("%Y-%m")


# ── 生命周期判定 ─────────────────────────────

def classify(first_seen: str, monthly: dict[str, int], today: dt.date) -> str:
    age = (today - dt.date.fromisoformat(first_seen)).days if first_seen else 0
    if monthly:
        seq = sorted(monthly.items())
        vals = [v for _, v in seq]
        if vals:
            peak_i = max(range(len(vals)), key=lambda i: vals[i])
            peak_v = vals[peak_i]
            recent3 = sum(vals[-3:]) / min(3, len(vals))
            trough = min(vals[peak_i + 1:] or vals[peak_i:])
            # 峰值距今月数
            y0, m0 = map(int, seq[peak_i][0].split("-"))
            months_since = (today.year - y0) * 12 + (today.month - m0)
            if (months_since >= 4 and peak_v >= 3 * max(trough, 1)
                    and recent3 >= 2 * max(trough, 1) and recent3 >= 0.25 * peak_v):
                return "二次爆火"
            base = vals[:12]
            if base and sum(base) / len(base) >= 1000 and age >= 120:
                return "长期大词"
    if age <= 90:
        return "新词"
    return "观察中"


# ── 读写 ─────────────────────────────

def load_library() -> dict[str, dict]:
    """返回 {词: record}；record 含 monthly={年月: 量}。文件不存在返回空库。"""
    if not LIBRARY.exists():
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(LIBRARY)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(h) if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    month_cols = [(i, h) for i, h in enumerate(header) if MONTH_RE.match(h)]
    lib: dict[str, dict] = {}
    for row in rows[1:]:
        word = str(row[idx["词"]] or "").strip().lower()
        if not word:
            continue
        monthly = {}
        for i, h in month_cols:
            if i < len(row) and isinstance(row[i], (int, float)):
                monthly[h] = int(row[i])
        lib[word] = {
            "word": word,
            "roots": str(row[idx["词根"]] or ""),
            "first": str(row[idx["首次出现"]] or ""),
            "last": str(row[idx["最近更新"]] or ""),
            "state": str(row[idx["状态"]] or ""),
            "gpts": row[idx["GPTs倍数"]] if idx.get("GPTs倍数") is not None else None,
            "peak": str(row[idx["峰值月"]] or ""),
            "note": str(row[idx["释义"]] or ""),
            "inner": row[idx["内页候选"]] if idx.get("内页候选") is not None else None,
            "dedicated": row[idx["专门页面"]] if idx.get("专门页面") is not None else None,
            "rec": str(row[idx["推荐日期"]] or ""),
            "remark": str(row[idx["备注"]] or ""),
            "monthly": monthly,
        }
    return lib


def save_library(lib: dict[str, dict], track_tag: str = "") -> str:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    months = sorted({m for r in lib.values() for m in r["monthly"]})
    headers = FIXED_HEADERS + months
    wb = Workbook()
    ws = wb.active
    ws.title = "词库"
    ws.append(headers + ([f"最近入库: {track_tag}"] if track_tag else []))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")

    state_fill = {"新词": "C6EFCE", "二次爆火": "FFD966", "长期大词": "BDD7EE",
                  "不值得做": "D9D9D9", "已推荐": "E2D5F1"}
    rows_sorted = sorted(
        lib.values(),
        key=lambda r: (STATE_ORDER.get(r["state"], 9),
                       -(r["gpts"] if isinstance(r["gpts"], (int, float)) else -1)))
    for r in rows_sorted:
        ws.append([r["word"], r["roots"], r["first"], r["last"], r["state"],
                   r["gpts"] if r["gpts"] is not None else "",
                   r["peak"], r["note"],
                   r["inner"] if r["inner"] is not None else "",
                   r["dedicated"] if r["dedicated"] is not None else "",
                   r["rec"], r["remark"]] + [r["monthly"].get(m, "") for m in months])
        fill = state_fill.get(r["state"])
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    for i, w in enumerate([36, 14, 11, 11, 10, 9, 9, 50, 9, 9, 11, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for j in range(len(FIXED_HEADERS) + 1, len(FIXED_HEADERS) + len(months) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 6
        col = get_column_letter(j)
        ws.conditional_formatting.add(
            f"{col}2:{col}{ws.max_row}",
            DataBarRule(start_type="num", start_value=0, end_type="max",
                        color="FFED7D31", showValue=False))
    ws.freeze_panes = "B2"
    LIBRARY.parent.mkdir(parents=True, exist_ok=True)
    wb.save(LIBRARY)
    return ws.title


def _to_int(v) -> int | None:
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


# ── 子命令 ─────────────────────────────

def cmd_merge(args) -> None:
    today = _today()
    lib = load_library()
    n_new = n_updated = 0

    rising = json.load(open(args.rising)) if args.rising else []
    for item in rising:
        word = str(item.get("query", "")).strip().lower()
        if not word:
            continue
        roots = ",".join(item.get("roots") or [])
        note = str(item.get("note", "")).strip()
        gpts = item.get("gpts")
        if word in lib:
            r = lib[word]
            n_updated += 1
            if roots and roots not in r["roots"]:
                r["roots"] = (r["roots"] + "," + roots).strip(",")
            r["last"] = today.isoformat()
            if gpts is not None:
                r["gpts"] = gpts
            if note and not r["note"]:
                r["note"] = note
        else:
            n_new += 1
            lib[word] = {"word": word, "roots": roots, "first": today.isoformat(),
                         "last": today.isoformat(), "state": "新词", "gpts": gpts,
                         "peak": "", "note": note, "inner": None, "dedicated": None,
                         "rec": "", "remark": "", "monthly": {}}
        lib[word]["word"] = word

    # 阶段 0.5 砍掉的词入库标「不值得做」（只新增，不降级库内已有词）
    n_dropped = 0
    if args.dropped:
        for item in json.load(open(args.dropped)):
            word = str(item.get("query", "")).strip().lower()
            if not word:
                continue
            reason = str(item.get("reason", "")).strip()
            if word in lib:
                if not lib[word]["remark"]:
                    lib[word]["remark"] = f"阶段0.5过滤: {reason}"
                continue
            n_dropped += 1
            lib[word] = {"word": word,
                         "roots": ",".join(item.get("roots") or []),
                         "first": today.isoformat(), "last": today.isoformat(),
                         "state": "不值得做", "gpts": None, "peak": "", "note": "",
                         "inner": None, "dedicated": None, "rec": "",
                         "remark": f"阶段0.5过滤: {reason}", "monthly": {}}

    # 月趋势数据 enrich（mine_round --export-monthly 产出）
    if args.monthly:
        data = json.load(open(args.monthly))
        n_monthly = 0
        for word, series in data.items():
            word = word.strip().lower()
            if word in lib:
                lib[word]["monthly"] = {k: int(v) for k, v in series.items()
                                        if _to_int(v) is not None}
                n_monthly += 1
        print(f"# 月趋势 enrich {n_monthly} 词", file=sys.stderr)

    # 状态重算（已推荐 粘性；不值得做 只有 monthly 显示真需求才升级）
    for word, r in lib.items():
        new_state = classify(r["first"], r["monthly"], today)
        if r["state"] == "已推荐":
            continue
        if r["state"] == "不值得做" and new_state not in ("二次爆火", "长期大词"):
            continue   # 粘性：除非月趋势证明是真需求，否则维持不值得做
        r["state"] = new_state
        if r["monthly"]:
            peak_ym = max(sorted(r["monthly"]), key=lambda m: r["monthly"][m])
            r["peak"] = peak_ym

    tag = f"{today.isoformat()} 新{n_new}/更新{n_updated}/过滤{n_dropped}"
    sheet = save_library(lib, tag)
    stats = {}
    for r in lib.values():
        stats[r["state"]] = stats.get(r["state"], 0) + 1
    print(json.dumps({"library": str(LIBRARY), "sheet": sheet,
                      "new": n_new, "updated": n_updated, "dropped": n_dropped,
                      "total": len(lib), "states": stats}, ensure_ascii=False))


def cmd_mark(args) -> None:
    lib = load_library()
    today = _today().isoformat()
    hit = 0
    for w in [x.strip().lower() for x in args.words.split(",") if x.strip()]:
        if w in lib:
            lib[w]["rec"] = today
            lib[w]["state"] = "已推荐"
            hit += 1
        else:
            print(f"# 词库中没有 {w}（先跑 merge 入库再 mark）", file=sys.stderr)
    save_library(lib)
    print(json.dumps({"marked": hit, "total": len(lib)}, ensure_ascii=False))


def cmd_stats(_args) -> None:
    lib = load_library()
    stats: dict[str, int] = {}
    for r in lib.values():
        stats[r["state"]] = stats.get(r["state"], 0) + 1
    print(json.dumps({"library": str(LIBRARY), "total": len(lib),
                      "states": stats}, ensure_ascii=False))


def cmd_export(args) -> None:
    """导出指定状态的词（默认 新词+二次爆火），格式同 rising_clean.json，可直接喂 mine_round。"""
    lib = load_library()
    states = set(args.states.split(",")) if args.states else {"新词", "二次爆火"}
    rows = [{"query": r["word"], "value": r["gpts"] or 0,
             "roots": [x for x in r["roots"].split(",") if x],
             "note": r["note"], "gpts": r["gpts"], "first_seen": r["first"]}
            for r in lib.values() if r["state"] in states]
    rows.sort(key=lambda x: -(x["gpts"] if isinstance(x["gpts"], (int, float)) else 0))
    out = json.dumps(rows, ensure_ascii=False, indent=1)
    if args.out:
        Path(args.out).write_text(out)
        print(f"# 已存 {args.out}（{len(rows)} 词）", file=sys.stderr)
    print(out)


def main() -> None:
    ap = argparse.ArgumentParser(description="持久词库管理（哥飞关键词库标准的本地版）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    m = sub.add_parser("merge", help="飙升词入库/更新状态（含 dropped 标注不值得做、monthly enrich）")
    m.add_argument("--rising", help="阶段0.5 过滤后的 rising_clean.json（可含 note/gpts 字段）")
    m.add_argument("--dropped", help="阶段0.5 过滤掉的 rising_dropped.json（含 reason 字段）")
    m.add_argument("--monthly", help="月趋势 {词: {年月: 量}} JSON（mine_round --export-monthly 产出）")
    m.set_defaults(fn=cmd_merge)

    k = sub.add_parser("mark", help="标记已推荐（业务拍板后）")
    k.add_argument("--words", required=True, help="逗号分隔的词")
    k.set_defaults(fn=cmd_mark)

    s = sub.add_parser("stats", help="词库统计")
    s.set_defaults(fn=cmd_stats)

    e = sub.add_parser("export", help="按状态导出词（默认 新词+二次爆火），可直接喂 mine_round --from-rising")
    e.add_argument("--states", help="逗号分隔的状态，默认 新词,二次爆火")
    e.add_argument("--out", help="另存文件")
    e.set_defaults(fn=cmd_export)

    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
